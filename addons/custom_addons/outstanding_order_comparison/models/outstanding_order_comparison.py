from odoo import fields, models
from odoo.exceptions import UserError


class OutstandingOrderComparison(models.Model):
    _name = 'outstanding.order.comparison'
    _description = 'Outstanding Order Comparison'

    name = fields.Char(string='Name', default='Outstanding Order Comparison')

    file_one = fields.Binary(string='Odoo Outstanding List', attachment=True)
    file_one_name = fields.Char(string='Odoo Outstanding List Name')

    file_two = fields.Binary(string='Higo Outstanding List', attachment=True)
    file_two_name = fields.Char(string='Higo Outstanding List Name')
    higo_outstanding_date = fields.Date(string='Higo Outstanding List Date')

    output_file = fields.Binary(string='Output File', attachment=True, readonly=True)
    output_file_name = fields.Char(string='Output File Name', readonly=True)

    def action_set_higo_outstanding_date(self):
        self.ensure_one()
        self.higo_outstanding_date = fields.Date.context_today(self)
        return {'type': 'ir.actions.client', 'tag': 'reload'}

    def action_run_program(self):
        """Alias for the Run button — delegates to action_process_files."""
        self.ensure_one()
        if not self.file_one or not self.file_two:
            raise UserError(
                'Please upload both the Odoo Outstanding List and the Higo Outstanding List before running.'
            )
        return self.action_process_files()

    def action_process_files(self):
        """
        Port of automatic_change_detection_OL_NVI20260403.py.

        file_one  → Odoo Outstanding List  (FILE_OLD / df_old)
        file_two  → Higo Outstanding List  (FILE_NEW / df_new)

        The date in the output filename is taken from higo_outstanding_date.
        """
        import base64
        import hashlib
        import io
        import re

        import numpy as np
        import pandas as pd
        from openpyxl import load_workbook as openpyxl_load_workbook
        from openpyxl.cell.rich_text import CellRichText, InlineFont, TextBlock
        from openpyxl.styles import Alignment, Border, Font as OpenpyxlFont, PatternFill, Side
        from openpyxl.utils import get_column_letter

        self.ensure_one()

        if not self.file_one or not self.file_two:
            raise UserError(
                'Please upload both the Odoo Outstanding List and the Higo Outstanding List.'
            )
        if not self.higo_outstanding_date:
            raise UserError(
                'Please fill in the Higo Outstanding List Date before running.'
            )

        # ------------------------------------------------------------------
        # DECODE FILES
        # ------------------------------------------------------------------
        odoo_bytes = base64.b64decode(self.file_one)   # FILE_OLD
        higo_bytes = base64.b64decode(self.file_two)   # FILE_NEW

        # ------------------------------------------------------------------
        # CONFIGURATION  (mirrors the original script)
        # ------------------------------------------------------------------
        KEY_COLUMN  = "KEY"
        KEY_FIELDS  = ['A&C PO', 'PO Line', 'PO Quantity']
        FILE_NEW_SHEETS = ["2026", "2025", "2024"]

        REQUESTED_OUTPUT_COLUMNS = [
            'Reference',
            'A&C PO',
            'PO Line',
            'Product',
            'A&C PN',
            'Higo PN',
            'shipping mode',
            'PO Quantity',
            'Unit price/pcs',
            'Delivery date requirment',
            'L/T',
            'Del. Date Confirmed',
        ]
        ITEM_COLUMN = 'Item'

        # Output filename — date from the Higo Outstanding List Date field
        date_str = self.higo_outstanding_date.strftime('%Y%m%d')
        output_filename = f"comparison_outstanding_orders_{date_str}.xlsx"

        # ------------------------------------------------------------------
        # HELPERS
        # ------------------------------------------------------------------
        def _normalise_sheet(df):
            df.columns = df.columns.str.strip()
            df.columns = df.columns.str.replace(r"\s+", " ", regex=True)
            df = df.replace(r'^\s*$', np.nan, regex=True)
            df = df.dropna(how='all')
            key_cols_present = [f for f in KEY_FIELDS if f in df.columns]
            if key_cols_present:
                df = df.dropna(subset=key_cols_present, how='all')
            return df

        def load_data(file_bytes, sheets=False):
            buf = io.BytesIO(file_bytes)
            xl = pd.ExcelFile(buf)

            if sheets is False:
                sheet_list = [xl.sheet_names[0]]
            elif sheets is None:
                sheet_list = xl.sheet_names
            else:
                available = set(xl.sheet_names)
                missing_sheets = [s for s in sheets if s not in available]
                if missing_sheets:
                    pass  # silently skip missing sheets (same as original)
                sheet_list = [s for s in sheets if s in available]
                if not sheet_list:
                    raise UserError(
                        f"None of the requested sheets {sheets} were found in the uploaded file."
                    )

            dfs = []
            for sheet in sheet_list:
                raw = pd.read_excel(xl, sheet_name=sheet)
                df = _normalise_sheet(raw)
                key_cols_present = [f for f in KEY_FIELDS if f in df.columns]
                if not key_cols_present:
                    continue
                dfs.append(df)

            if not dfs:
                raise UserError("No usable sheets (containing key columns) were found in the uploaded file.")

            return pd.concat(dfs, ignore_index=True)

        def resolve_columns(existing_columns, requested_columns):
            normalized_lookup = {
                re.sub(r"\s+", " ", str(col).strip()).casefold(): col
                for col in existing_columns
            }
            resolved, missing = [], []
            for col in requested_columns:
                key = re.sub(r"\s+", " ", str(col).strip()).casefold()
                match = normalized_lookup.get(key)
                if match is not None:
                    resolved.append(match)
                else:
                    missing.append(col)
            return resolved, missing

        def make_key(row):
            def clean_value(value):
                if pd.isna(value):
                    return ""
                if isinstance(value, float) and value.is_integer():
                    return str(int(value))
                return str(value).strip()
            key_string = "|".join(clean_value(row[field]) for field in KEY_FIELDS)
            return hashlib.sha256(key_string.encode("utf-8")).hexdigest()

        def validate(df1, df2, key):
            if key not in df1.columns or key not in df2.columns:
                raise UserError(f"Key column '{key}' missing in one of the files.")
            dup1 = df1[df1[key].duplicated(keep=False)]
            if not dup1.empty:
                sample = dup1.head(10)[KEY_FIELDS + [key]] if all(f in dup1.columns for f in KEY_FIELDS) else dup1.head(10)
                raise UserError(
                    f"Duplicate keys in the Higo file ({len(dup1)} rows).\n"
                    f"Sample:\n{sample.to_string(index=False)}"
                )
            dup2 = df2[df2[key].duplicated(keep=False)]
            if not dup2.empty:
                sample = dup2.head(10)[KEY_FIELDS + [key]] if all(f in dup2.columns for f in KEY_FIELDS) else dup2.head(10)
                raise UserError(
                    f"Duplicate keys in the Odoo file ({len(dup2)} rows).\n"
                    f"Sample:\n{sample.to_string(index=False)}"
                )

        def _normalise_for_cmp(val):
            if pd.isna(val):
                return val
            if isinstance(val, str):
                return val.strip().casefold()
            return val

        # ------------------------------------------------------------------
        # LOAD DATA
        # ------------------------------------------------------------------
        df_new = load_data(higo_bytes, sheets=FILE_NEW_SHEETS)   # Higo = FILE_NEW
        df_old = load_data(odoo_bytes, sheets=False)              # Odoo  = FILE_OLD

        # ------------------------------------------------------------------
        # KEY GENERATION
        # ------------------------------------------------------------------
        df_new[KEY_COLUMN] = df_new.apply(make_key, axis=1)
        df_old[KEY_COLUMN] = df_old.apply(make_key, axis=1)

        # ------------------------------------------------------------------
        # VALIDATION
        # ------------------------------------------------------------------
        validate(df_new, df_old, KEY_COLUMN)

        # ------------------------------------------------------------------
        # COLUMN MEMBERSHIP SETS
        # ------------------------------------------------------------------
        _new_only_cols = set(df_new.columns) - set(df_old.columns) - {KEY_COLUMN}
        _old_only_cols = set(df_old.columns) - set(df_new.columns) - {KEY_COLUMN}

        # ------------------------------------------------------------------
        # MERGE
        # ------------------------------------------------------------------
        merged = df_new.merge(
            df_old,
            on=KEY_COLUMN,
            how="outer",
            suffixes=("_new", "_old"),
            indicator=True,
        )

        # ------------------------------------------------------------------
        # ROW STATUS
        # ------------------------------------------------------------------
        merged["row_status"] = merged["_merge"].apply(
            lambda x: "NEW" if x == "left_only" else "DELETED" if x == "right_only" else "EXISTING"
        )

        # ------------------------------------------------------------------
        # COLUMN RESOLUTION
        # ------------------------------------------------------------------
        _all_columns = list(df_new.columns) + [c for c in df_old.columns if c not in set(df_new.columns)]
        resolved_compare_columns, missing_compare_columns = resolve_columns(_all_columns, REQUESTED_OUTPUT_COLUMNS)
        resolved_item_columns, _ = resolve_columns(df_new.columns, [ITEM_COLUMN])
        item_column_name = resolved_item_columns[0] if resolved_item_columns else None
        display_columns = ([item_column_name] if item_column_name else []) + resolved_compare_columns

        _item_col_set = {item_column_name} if item_column_name else set()
        compare_columns_fields = [
            col for col in resolved_compare_columns
            if col not in _old_only_cols and col not in _item_col_set
        ]

        # ------------------------------------------------------------------
        # CHANGE DETECTION
        # ------------------------------------------------------------------
        def compare_columns(row):
            changes = []
            for col in compare_columns_fields:
                if col in _new_only_cols:
                    new_val, old_val = row.get(col, np.nan), np.nan
                elif col in _old_only_cols:
                    new_val, old_val = np.nan, row.get(col, np.nan)
                else:
                    new_val = row.get(f"{col}_new", np.nan)
                    old_val = row.get(f"{col}_old", np.nan)
                new_cmp = _normalise_for_cmp(new_val)
                old_cmp = _normalise_for_cmp(old_val)
                if pd.isna(new_cmp) and pd.isna(old_cmp):
                    continue
                if new_cmp != old_cmp:
                    changes.append(f"{col}: '{old_val}' → '{new_val}'")
            return "; ".join(changes)

        merged["changes"] = merged.apply(compare_columns, axis=1)

        # ------------------------------------------------------------------
        # FINAL STATUS
        # ------------------------------------------------------------------
        def finalize_status(row):
            if row["row_status"] in ["NEW", "DELETED"]:
                return row["row_status"]
            return "MODIFIED" if row["changes"] else "UNCHANGED"

        merged["final_status"] = merged.apply(finalize_status, axis=1)

        # ------------------------------------------------------------------
        # CLEAN OUTPUT DATAFRAME
        # ------------------------------------------------------------------
        valid_output_cols = []
        rename_map = {}
        for col in display_columns:
            if col in _new_only_cols or col in _old_only_cols:
                if col in merged.columns:
                    valid_output_cols.append(col)
            else:
                merged_col = f"{col}_new"
                if merged_col in merged.columns:
                    valid_output_cols.append(merged_col)
                    rename_map[merged_col] = col

        output_columns = ["final_status", "changes"] + valid_output_cols
        final_df = merged[output_columns].copy()
        final_df = final_df.rename(columns=rename_map)

        _group_map = {"DELETED": 0, "MODIFIED": 1, "UNCHANGED": 1, "NEW": 2}
        final_df["_sort_group"] = final_df["final_status"].map(_group_map)
        if item_column_name and item_column_name in final_df.columns:
            final_df[item_column_name] = pd.to_numeric(final_df[item_column_name], errors="coerce")
            final_df = final_df.sort_values(
                ["_sort_group", item_column_name], ascending=[True, True], na_position="last"
            )
        else:
            final_df = final_df.sort_values("_sort_group", ascending=True)
        final_df = final_df.drop(columns=["_sort_group"])

        # ------------------------------------------------------------------
        # BUILD EXCEL OUTPUT IN MEMORY
        # ------------------------------------------------------------------
        output_buf = io.BytesIO()

        with pd.ExcelWriter(output_buf, engine="openpyxl") as writer:

            # ---- Fill colours (identical to original) ----
            FILL_REMOVED   = PatternFill(start_color="CCFFCC", end_color="CCFFCC", fill_type="solid")
            FILL_ADDED     = PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid")
            FILL_CHANGED_A = PatternFill(start_color="FFFF99", end_color="FFFF99", fill_type="solid")
            FILL_CHANGED_B = PatternFill(start_color="FFE033", end_color="FFE033", fill_type="solid")
            FILL_PREV_A    = PatternFill(start_color="F0F0F0", end_color="F0F0F0", fill_type="solid")
            FILL_PREV_B    = PatternFill(start_color="D8D8D8", end_color="D8D8D8", fill_type="solid")
            FILL_HEADER    = PatternFill(start_color="1F7A4A", end_color="1F7A4A", fill_type="solid")

            workbook = writer.book
            ws_vc = workbook.create_sheet("Visual Comparison")

            vis_columns = display_columns
            header_vc   = ["Row Type"] + vis_columns
            ws_vc.append(header_vc)
            for cell in ws_vc[1]:
                cell.font = OpenpyxlFont(bold=True, color="FFFFFF")
                cell.fill = FILL_HEADER

            def get_cell_val(series, col):
                try:
                    val = series[col]
                    return "" if pd.isna(val) else val
                except (KeyError, TypeError, ValueError):
                    return ""

            def cells_differ(v1, v2):
                try:
                    v1_na = pd.isna(v1)
                    v2_na = pd.isna(v2)
                except (TypeError, ValueError):
                    v1_na, v2_na = False, False
                if v1_na and v2_na:
                    return False
                if v1_na != v2_na:
                    return True
                if isinstance(v1, str) and isinstance(v2, str):
                    return v1.strip().casefold() != v2.strip().casefold()
                try:
                    return bool(v1 != v2)
                except Exception:
                    return str(v1) != str(v2)

            def write_vis_row(ws, row_type, series, fill_row=None, highlight_cols=None, fill_changed=None):
                highlight_cols = highlight_cols or set()
                fill_changed   = fill_changed or FILL_CHANGED_A
                data = [row_type] + [get_cell_val(series, col) for col in vis_columns]
                ws.append(data)
                er = ws.max_row
                for col_idx, col_name in enumerate(header_vc, start=1):
                    cell = ws.cell(row=er, column=col_idx)
                    if col_name in highlight_cols:
                        cell.fill = fill_changed
                    elif fill_row is not None:
                        cell.fill = fill_row

            # 1. DELETED rows (IN ODOO & NOT ON HIGO LIST) — light green
            for _, mrow in merged[merged["final_status"] == "DELETED"].iterrows():
                key = mrow[KEY_COLUMN]
                old_match = df_old[df_old[KEY_COLUMN] == key]
                if not old_match.empty:
                    write_vis_row(ws_vc, "IN ODOO & NOT ON HIGO LIST", old_match.iloc[0], fill_row=FILL_REMOVED)

            # 2. MODIFIED + UNCHANGED (sorted by Item#), then NEW
            remaining_vc = merged[merged["final_status"] != "DELETED"].copy()
            middle_vc    = remaining_vc[remaining_vc["final_status"].isin(["MODIFIED", "UNCHANGED"])].copy()
            new_vc       = remaining_vc[remaining_vc["final_status"] == "NEW"].copy()

            if item_column_name and item_column_name in middle_vc.columns:
                middle_vc["_sort_item"] = pd.to_numeric(middle_vc[item_column_name], errors="coerce")
                middle_vc = middle_vc.sort_values("_sort_item", ascending=True, na_position="last")

            _modified_pair_idx = 0

            for _, mrow in list(middle_vc.iterrows()) + list(new_vc.iterrows()):
                status = str(mrow["final_status"])
                key    = mrow[KEY_COLUMN]

                if status == "NEW":
                    new_match = df_new[df_new[KEY_COLUMN] == key]
                    if not new_match.empty:
                        write_vis_row(ws_vc, "NOT IN ODOO & ON HIGO LIST", new_match.iloc[0], fill_row=FILL_ADDED)

                elif status == "UNCHANGED":
                    new_match = df_new[df_new[KEY_COLUMN] == key]
                    if not new_match.empty:
                        write_vis_row(ws_vc, "", new_match.iloc[0])

                elif status == "MODIFIED":
                    new_match = df_new[df_new[KEY_COLUMN] == key]
                    old_match = df_old[df_old[KEY_COLUMN] == key]
                    if new_match.empty or old_match.empty:
                        continue
                    new_series = new_match.iloc[0]
                    old_series = old_match.iloc[0]

                    _use_a    = (_modified_pair_idx % 2 == 0)
                    fill_chg  = FILL_CHANGED_A if _use_a else FILL_CHANGED_B
                    fill_prev = FILL_PREV_A    if _use_a else FILL_PREV_B
                    _modified_pair_idx += 1

                    changed_cols = {
                        col for col in vis_columns
                        if col not in _old_only_cols
                        and col not in _item_col_set
                        and cells_differ(
                            new_series.get(col, np.nan),
                            old_series.get(col, np.nan),
                        )
                    }

                    # Current (Higo) row
                    write_vis_row(ws_vc, "HIGO LIST", new_series, highlight_cols=changed_cols, fill_changed=fill_chg)

                    # Previous (Odoo) row — grey background, same yellow highlights on changed cells
                    prev_data = ["ODOO"] + [get_cell_val(old_series, col) for col in vis_columns]
                    ws_vc.append(prev_data)
                    prev_er = ws_vc.max_row
                    for col_idx, col_name in enumerate(header_vc, start=1):
                        cell = ws_vc.cell(row=prev_er, column=col_idx)
                        if col_name in changed_cols:
                            cell.fill = fill_chg
                        else:
                            cell.fill = fill_prev

            # ---- Post-processing: borders, alignment, row heights ----
            thin       = Side(border_style="thin", color="000000")
            all_border = Border(left=thin, right=thin, top=thin, bottom=thin)
            center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

            for row_idx, row in enumerate(ws_vc.iter_rows(), start=1):
                ws_vc.row_dimensions[row_idx].height = 30
                for cell in row:
                    cell.border    = all_border
                    cell.alignment = center_align

            # ---- Column widths from the Higo (source) file ----
            _src_wb = openpyxl_load_workbook(io.BytesIO(higo_bytes), data_only=True)
            _src_ws = _src_wb.active
            _src_col_widths = {}
            for cell in next(_src_ws.iter_rows(min_row=1, max_row=1)):
                if cell.value is not None:
                    cd = _src_ws.column_dimensions.get(get_column_letter(cell.column))
                    if cd and cd.width and cd.width > 1:
                        _src_col_widths[str(cell.value).strip()] = cd.width
            _src_wb.close()

            for col_idx, col_name in enumerate(header_vc, start=1):
                col_letter = get_column_letter(col_idx)
                if col_name in _src_col_widths:
                    ws_vc.column_dimensions[col_letter].width = _src_col_widths[col_name]
                else:
                    max_len = 0
                    for row in ws_vc.iter_rows(min_col=col_idx, max_col=col_idx):
                        for cell in row:
                            if cell.value is not None:
                                longest_line = max(
                                    (len(line) for line in str(cell.value).splitlines()), default=0
                                )
                                max_len = max(max_len, longest_line)
                    ws_vc.column_dimensions[col_letter].width = max(max_len + 4, 10)

            # ---- Fixed-width overrides (same as original) ----
            _fixed_width_cols = {"Delivery date requirment", "L/T", "PO Issue Date"}
            _fixed_width_map  = {"Product": 66.43, "Row Type": 17.86, "Del. Date Confirmed": 17.86}
            for col_idx, col_name in enumerate(header_vc, start=1):
                if col_name in _fixed_width_cols:
                    ws_vc.column_dimensions[get_column_letter(col_idx)].width = 17.86
                elif col_name in _fixed_width_map:
                    ws_vc.column_dimensions[get_column_letter(col_idx)].width = _fixed_width_map[col_name]

            # ---- Bold "NOT" in Row Type column ----
            _bold_inline   = InlineFont(b=True)
            _normal_inline = InlineFont()
            for _r in range(2, ws_vc.max_row + 1):
                _cell = ws_vc.cell(row=_r, column=1)
                if isinstance(_cell.value, str) and "NOT" in _cell.value:
                    _parts = _cell.value.split("NOT")
                    _rich  = []
                    for _i, _p in enumerate(_parts):
                        if _p:
                            _rich.append(TextBlock(_normal_inline, _p))
                        if _i < len(_parts) - 1:
                            _rich.append(TextBlock(_bold_inline, "NOT"))
                    _cell.value = CellRichText(*_rich)

            # ---- Del. Date Confirmed → ✓ / ✗ ----
            _dc_col = "Del. Date Confirmed"
            if _dc_col in header_vc:
                _dc_col_idx = header_vc.index(_dc_col) + 1
                for _r in range(2, ws_vc.max_row + 1):
                    _cell = ws_vc.cell(row=_r, column=_dc_col_idx)
                    _val  = _cell.value
                    if _val is None or str(_val).strip() in ("", "NaT", "nan", "None"):
                        _cell.value = "✗"
                    else:
                        _cell.value = "✓"

            # ---- Freeze top row ----
            ws_vc.freeze_panes = "A2"

        # ------------------------------------------------------------------
        # STORE OUTPUT ON THE RECORD
        # ------------------------------------------------------------------
        output_bytes = output_buf.getvalue()
        self.write({
            'output_file':      base64.b64encode(output_bytes),
            'output_file_name': output_filename,
        })

        download_url = (
            f'/web/content?model={self._name}'
            f'&id={self.id}'
            f'&field=output_file'
            f'&filename={output_filename}'
            f'&download=true'
        )

        return {
            'type': 'ir.actions.act_url',
            'url': download_url,
            'target': 'new',
        }
